"""
Génère tests/reference_excel.json : valeurs de référence calculées par l'Excel d'origine.

Principe : pour chaque scénario (valeurs INIES collées + DVR), on écrit les entrées dans
l'onglet « Interface +A2 » d'une copie du classeur, on fait recalculer la copie par
LibreOffice, puis on relit les résultats (cellules D23:G24 de l'onglet).

Usage (depuis la racine du projet) :  python3 tests/generer_reference_excel.py
Prérequis : openpyxl et LibreOffice (commande `soffice`).
"""
import json
import shutil
import subprocess
import tempfile
from pathlib import Path

import openpyxl

# ----------------------------------------------------------------------------
# VARIABLES (tout est regroupé ici)
# ----------------------------------------------------------------------------
RACINE = Path(__file__).resolve().parent.parent
EXCEL_ORIGINE = RACINE / "reference" / "ACV_Dynamique_INIES_V2.xlsx"
SORTIE_JSON = RACINE / "tests" / "reference_excel.json"
ONGLET = "Interface +A2"

COLONNES_VALEURS = "CDEFGHIJKLMNO"      # 13 colonnes collées dans la zone rose (ligne 10)
LIGNE_TEXTE = 10                        # ligne où l'utilisateur colle (texte "1.54e+1")
LIGNE_CHIFFRES = 15                     # ligne « Chiffres » (nombres) que l'on écrase
CELLULE_DVR = "C4"
# Cellules de résultat lues : (clé JSON, cellule)
CELLULES_RESULTAT = {
    "dyn_t0": "D24", "dyn_milieu": "E24", "dyn_t50": "F24", "dyn_total": "G24",
    "stat_t0": "D23", "stat_milieu": "E23", "stat_t50": "F23", "stat_total": "G23",
}

# Valeurs collées (texte tel que copié depuis INIES), dans l'ordre des 13 colonnes
CLT_MATHIS = "-6.04e+2 1.38e+1 0.00e+0 0.00e+0 0.00e+0 0.00e+0 0.00e+0 0.00e+0 0.00e+0 0.00e+0 7.58e+2 1.67e+2 -2.33e+2"
FICHE_46517 = "1.54e+1 5.14e+0 0.00e+0 5.95e-3 0.00e+0 0.00e+0 0.00e+0 0.00e+0 0.00e+0 5.95e-3 4.23e-1 2.10e+1 -9.65e-2"
# Cas avec phase d'utilisation importante (pour tester la division par la DVR)
AVEC_UTILISATION = "1.20e+1 3.00e+0 0.00e+0 4.00e+0 0.00e+0 2.00e+0 0.00e+0 1.50e+1 0.00e+0 2.10e+1 5.00e+0 4.10e+1 -3.00e+0"

SCENARIOS = [
    ("CLT Mathis, DVR 100 (cas de l'Excel)", CLT_MATHIS, 100),
    ("CLT Mathis, DVR 30", CLT_MATHIS, 30),
    ("Fiche 46517, DVR 50", FICHE_46517, 50),
    ("Fiche 46517, DVR 60", FICHE_46517, 60),
    ("Fiche 46517, DVR 30 (reliquat 2/3)", FICHE_46517, 30),
    ("Fiche 46517, DVR 25 (reliquat nul)", FICHE_46517, 25),
    ("Fiche 46517, DVR 20 (reliquat 1/2)", FICHE_46517, 20),
    ("Fiche 46517, DVR 10 (reliquat nul)", FICHE_46517, 10),
    ("Fiche 46517, DVR 7", FICHE_46517, 7),
    ("Fiche 46517, DVR 1", FICHE_46517, 1),
    ("Utilisation importante, DVR 40", AVEC_UTILISATION, 40),
    ("Utilisation importante, DVR 100", AVEC_UTILISATION, 100),
]


def preparer_copie(dossier: Path, indice: int, texte: str, dvr: int) -> Path:
    """Écrit les entrées du scénario dans une copie du classeur."""
    wb = openpyxl.load_workbook(EXCEL_ORIGINE)
    ws = wb[ONGLET]
    ws[CELLULE_DVR] = dvr
    for col, jeton in zip(COLONNES_VALEURS, texte.split()):
        ws[f"{col}{LIGNE_TEXTE}"] = jeton                       # texte brut collé
        ws[f"{col}{LIGNE_CHIFFRES}"] = float(jeton)             # évite la dépendance à la locale de VALUE()
    chemin = dossier / f"scenario_{indice}.xlsx"
    wb.save(chemin)
    return chemin


def recalculer(chemin: Path, dossier_sortie: Path) -> Path:
    """Fait recalculer le classeur par LibreOffice (conversion xlsx -> xlsx)."""
    dossier_sortie.mkdir(exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(dossier_sortie), str(chemin)],
        check=True, capture_output=True, timeout=180,
    )
    return dossier_sortie / chemin.name


def main() -> None:
    reference = []
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        for i, (nom, texte, dvr) in enumerate(SCENARIOS):
            calcule = recalculer(preparer_copie(tmp, i, texte, dvr), tmp / "calcule")
            ws = openpyxl.load_workbook(calcule, data_only=True)[ONGLET]
            attendu = {cle: ws[cell].value for cle, cell in CELLULES_RESULTAT.items()}
            reference.append({"nom": nom, "texte_inies": texte, "dvr": dvr, "attendu": attendu})
            print(f"{nom:45s} dyn_total={attendu['dyn_total']!r:>22}  stat_total={attendu['stat_total']!r}")
    SORTIE_JSON.write_text(json.dumps(reference, ensure_ascii=False, indent=2), encoding="utf-8")
    print("écrit :", SORTIE_JSON)


if __name__ == "__main__":
    main()
